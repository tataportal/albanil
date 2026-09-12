"""Read the client workbook; export only customer-facing catalog data, never store prices.
Run with bundled Python: import-consolidated.py [workbook.xlsx]
"""
from pathlib import Path
from collections import Counter
from datetime import datetime, timezone
import json, sys, re, unicodedata
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
source=Path(sys.argv[1]) if len(sys.argv)>1 else ROOT.parent/'Consolidado_Productos_Albanil.xlsx'
values=openpyxl.load_workbook(source,read_only=True,data_only=True)['Productos']
formats=openpyxl.load_workbook(source,read_only=True,data_only=False)['Productos']
original=json.loads((ROOT/'docs/catalog.json').read_text())
original_ids={p['id'] for p in original}
rows=[(i,list(row)) for i,row in enumerate(values.iter_rows(min_row=2,values_only=True),2) if any(v is not None for v in row)]
counts=Counter(r[0] for _,r in rows if r[0] is not None)
map_path=ROOT/'scripts/catalog-id-map.json'
registry=json.loads(map_path.read_text()) if map_path.exists() else {}
used=original_ids | {r[0] for _,r in rows if r[0] is not None} | {r['id'] for r in registry.values()}
normalize=lambda s: re.sub(r'\s+',' ',unicodedata.normalize('NFKC',str(s or '')).strip()).casefold()
occurrences=Counter();out=[]
for index,row in rows:
    source_id,title,category,brand,price,store_price,state,stock=row
    title=' '.join(title.split())
    base_key='|'.join(map(normalize,[source_id,title,category,brand]))
    occurrences[base_key]+=1
    key=base_key+'|'+str(occurrences[base_key])
    # ID 340 is the tapón in the existing site. The copper rod is a distinct item.
    duplicate=source_id==340 and 'varilla' in title.lower()
    if key not in registry:
        if source_id is not None and not duplicate:
            registry[key]={'id':int(source_id),'reference':str(source_id)}
        else:
            candidate=max([100000]+[i for i in used if i>=100000])+1
            used.add(candidate)
            registry[key]={'id':candidate,'reference':'340-2' if duplicate else f'N-{candidate-100000:04d}'}
    identity=registry[key]
    web_format=formats.cell(index,5).number_format
    if '[$$' in web_format:currency='USD'
    elif 'S/.' in web_format:currency='PEN'
    else:raise ValueError(f'Unknown currency format at row {index}: {web_format}')
    if isinstance(price,str) and price.strip()=='CONSULTAR PRECIO A NUESTROS ASESORES':price=None
    if price is not None and (not isinstance(price,(int,float)) or price<=0):raise ValueError(f'Invalid price at row {index}')
    if stock is not None and (not isinstance(stock,(int,float)) or stock<0):raise ValueError(f'Invalid stock at row {index}')
    out.append({**identity,'sourceId':source_id,'sourceRow':index,'title':title,'brand':brand,'category':category,'price':price,'currency':currency,'stock':stock,'state':state,'availability':'Por confirmar' if stock is None else 'Disponible' if stock>0 else 'Agotado','unit':'','tax':'confirmar'})
assert len({p['id'] for p in out})==len(out)
assert len({p['reference'] for p in out})==len(out)
map_path.write_text(json.dumps(registry,ensure_ascii=False,indent=2)+'\n')
(ROOT/'scripts/consolidated-products.json').write_text(json.dumps({'source':source.name,'importedAt':datetime.now(timezone.utc).isoformat(),'products':out},ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'rows':len(out),'currencies':dict(Counter(p['currency'] for p in out)),'existing_ids':sum(p['id'] in original_ids for p in out),'new_products':sum(p['id'] not in original_ids for p in out),'missing_stock':sum(p['stock'] is None for p in out),'duplicate_reference':[p['reference'] for p in out if p['reference']=='340-2']},ensure_ascii=False))
